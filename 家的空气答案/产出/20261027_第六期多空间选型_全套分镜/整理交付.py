from pathlib import Path
import json, shutil, html, zipfile
from PIL import Image
import openpyxl

root=Path(__file__).parent
source=Path('/Users/weipeilong/Desktop/璀璨中央空调/璀璨日常整理留存/10月/视频/【家的空气答案】第六期 多空间选型脚本.xlsx')
sheet=openpyxl.load_workbook(source,data_only=True)['Sheet1']
records=json.loads((root/'生成记录.json').read_text())
cards=[]
md=['# 第六期 · 多空间选型分镜图','', '对应本期脚本14个主镜头，9:16竖版静态预览。使用内置图像生成工具，产品镜头输入官方清晰原图；不是最终视频或安装施工图。','', '原始脚本：'+str(source),'', '|镜号|时间|画面|口播|','|---|---|---|---|']
prompts=[]
for r in records:
    i=int(r['id']); row=i+4
    dest=root/'分镜图'/f"{r['id']}_{r['title']}.png"
    if not dest.exists(): shutil.copy2(r['source'],dest)
    with Image.open(dest) as im:
        im.verify()
    with Image.open(dest) as im:
        dims=im.size
    r['file']=str(dest.relative_to(root));r['size']=dims
    time=str(sheet.cell(row,1).value or '')
    vo=str(sheet.cell(row,4).value or '')
    r['time']=time;r['vo']=vo
    md.append(f"|{r['id']}|{time}|{r['title']}|{vo.replace(chr(10),'<br>')}|")
    cards.append(f'''<article><a href="{html.escape(r['file'])}" target="_blank"><img src="{html.escape(r['file'])}" alt="{html.escape(r['title'])}"></a><div class="body"><span class="meta">{r['id']} / {time} / {dims[0]} × {dims[1]}</span><h2>{html.escape(r['title'])}</h2><p>{html.escape(vo)}</p></div></article>''')
    prompts.append(f"## {r['id']} {r['title']}\n\n{r['prompt']}\n")
(root/'分镜对应表.md').write_text('\n'.join(md),encoding='utf-8')
(root/'生成提示词.md').write_text('# 完整生成提示词\n\n工具：内置 image_gen。产品图以项目官方资产为输入参考。\n\n'+'\n'.join(prompts),encoding='utf-8')
(root/'分镜总览.html').write_text('''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>家的空气答案 · 第六期全套分镜</title><style>*{box-sizing:border-box}body{margin:0;background:#eeeae2;color:#332c26;font:15px -apple-system,BlinkMacSystemFont,"PingFang SC",sans-serif}header{padding:42px 4vw 28px}header .eyebrow{letter-spacing:.2em;font-size:12px;color:#8b7865}h1{font-size:32px;letter-spacing:.025em;margin:12px 0}header p{color:#766a5f}main{padding:0 4vw 48px;display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:24px}article{background:#faf8f3;border-radius:4px;overflow:hidden;box-shadow:0 5px 22px #332c2610}img{display:block;width:100%;height:auto}a{color:inherit}.body{padding:18px}.meta{font-size:11px;color:#8b7865}h2{font-size:17px;margin:8px 0 12px}p{line-height:1.8;margin:0;font-size:13px}footer{padding:20px 4vw 50px;color:#8b7865;font-size:12px}@media(max-width:1100px){main{grid-template-columns:repeat(3,1fr)}}@media(max-width:760px){main{grid-template-columns:repeat(2,1fr);gap:12px}.body{padding:12px}h1{font-size:26px}}@media(max-width:440px){main{grid-template-columns:1fr}}</style><header><div class="eyebrow">家的空气答案 / EPISODE 06</div><h1>多空间选型 · 全套分镜</h1><p>14个主镜头 · 约84秒 · 9:16竖版 · 点击图片查看原图</p></header><main>'''+''.join(cards)+'''</main><footer>当前为静态分镜预览，气流、管路、净化为视觉示意；片内文字为构图预览，正式视频排版另行制作。</footer></html>''',encoding='utf-8')
(root/'生成记录.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
if len(records)==14:
    zip_path=root/'第六期_全套14镜分镜图.zip'
    with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted((root/'分镜图').glob('*.png')): z.write(p,p.relative_to(root))
        for n in ['分镜总览.html','分镜对应表.md','生成提示词.md']: z.write(root/n,n)
    print('ZIP',zip_path,zip_path.stat().st_size)
print('完成',len(records),'张；尺寸',sorted({tuple(r['size']) for r in records}))
